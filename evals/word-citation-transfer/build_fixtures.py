from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree as E
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json,hashlib
OUT=Path(__file__).resolve().parent/'fixtures'
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
B='http://schemas.openxmlformats.org/officeDocument/2006/bibliography'
R='http://schemas.openxmlformats.org/package/2006/relationships'
REL='http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
CT='http://schemas.openxmlformats.org/package/2006/content-types'
DS='http://schemas.openxmlformats.org/officeDocument/2006/customXml'
def x(tag):return OxmlElement('w:'+tag)
def field(parent,code,cache):
 for kind,text in [('begin',None),('instruction',code),('separate',None),('text',cache),('end',None)]:
  r=x('r')
  if kind=='instruction':
   el=x('instrText');el.set(qn('xml:space'),'preserve');el.text=' '+text+' '
  elif kind=='text':el=x('t');el.text=text
  else:el=x('fldChar');el.set(qn('w:fldCharType'),kind)
  r.append(el);parent.append(r)
def text(parent,s):
 r=x('r');t=x('t');t.set(qn('xml:space'),'preserve');t.text=s;r.append(t);parent.append(r)
def xml(root):return E.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
def build(name,clean=False):
 cfg=([('ValeA','Tamsin','Vale','2024','Alpha Sensor Calibration'),('ValeB','Tamsin','Vale','2024','Beta Sensor Calibration'),('NgNote','Rui','Ng','2022','Quiet Instrument Records'),('ImaniCaption','Ada','Imani','2021','Diagram Conventions'),('MarinUnused','Lou','Marin','2020','Spare Workshop Notes'),('SatoReading','Mei','Sato','2019','Further Instrument Reading'),('LumenLibrary','Kai','Lumen','2018','Stored but Unlisted Records')] if not clean else [('EllisA','Morgan','Ellis','2020','Archive Register Methods'),('EllisB','Morgan','Ellis','2020','Museum Register Methods'),('PazNote','Lina','Paz','2018','Oral Catalogue Notes'),('KlineCaption','Evan','Kline','2022','Collection Diagrams'),('HuangReading','Ren','Huang','2017','Further Curatorial Reading'),('DaneLibrary','Jo','Dane','2016','Stored Collection Index')])
 src=[]
 for i,(tag,first,last,year,title) in enumerate(cfg,1):src.append(dict(tag=tag,first=first,last=last,year=year,title=title,guid='{'+f'{2 if clean else 1}{i:07d}-1111-4111-8111-{i:012d}'+'}'))
 def cite(i,suffix=''):return '('+src[i]['last']+', '+src[i]['year']+suffix+')'
 def entry(i,suffix=''):
  s=src[i];return f"{s['last']}, {s['first'][0]}. ({s['year']}{suffix}). {s['title']}. Example Press."
 d=Document();sec=d.sections[0];sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.8)
 for sn in ['Normal','Title','Heading 1','Caption']:
  st=d.styles[sn];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
 d.styles['Normal'].font.size=Pt(11);d.styles['Normal'].paragraph_format.space_after=Pt(6);d.styles['Title'].font.size=Pt(19)
 for b in d.styles.element.xpath('.//w:pBdr'):b.getparent().remove(b)
 d.core_properties.author=d.core_properties.last_modified_by='Fictional Test Author'
 d.add_paragraph('Collection record methods' if clean else 'Instrument record methods','Title')
 d.add_paragraph('This constructed document contains fictional references for a citation audit. It is not a research submission.')
 p=d.add_paragraph('The catalogue containss two registers ' if clean else 'The alpha calibration protocol is recorded in ')
 field(p._p,'CITATION '+src[0]['tag']+r' \l 1033',cite(0,'a' if clean else 'b'));p.add_run('.')
 p=d.add_paragraph('The second register is documented separately ' if clean else 'The beta calibration protocol has a separate record ')
 field(p._p,'CITATION '+src[1]['tag']+r' \l 1033',cite(1,'b' if clean else ''));p.add_run('.')
 p=d.add_paragraph('An explanatory source appears only in the attached note')
 r=p.add_run()._r;note=x('footnoteReference');note.set(qn('w:id'),'1');r.append(note);p.add_run('.')
 d.add_paragraph('Recorded items → Organized catalogue → Comparison','Normal')
 p=d.add_paragraph('Figure ','Caption');bm=x('bookmarkStart');bm.set(qn('w:id'),'19');bm.set(qn('w:name'),'figProcess');p._p.append(bm)
 field(p._p,r'SEQ Figure \* ARABIC','1');end=x('bookmarkEnd');end.set(qn('w:id'),'19');p._p.append(end)
 p.add_run('  Overview of the record process, based on ');field(p._p,'CITATION '+src[3]['tag']+r' \l 1033',cite(3));p.add_run('.')
 p=d.add_paragraph('The process is summarized in Figure ');field(p._p,r'REF figProcess \h','1');p.add_run('.')
 d.add_paragraph('References','Heading 1')
 entries=[entry(0,'a'),entry(1,'b'),entry(2),entry(3)]+([] if clean else [entry(4)])
 # The intentionally authored caches are test inputs, not claimed native Word refresh outputs.
 p=d.add_paragraph();field(p._p,r'BIBLIOGRAPHY \l 1033','\n'.join(entries))
 # Convert only test-input cache newlines into line breaks for readable rendering.
 for el in list(p._p.iter(qn('w:t'))):
  if '\n' in (el.text or ''):
   chunks=el.text.split('\n');el.text=chunks[0];r=el.getparent()
   for chunk in chunks[1:]:r.append(x('br'));t=x('t');t.text=chunk;r.append(t)
 d.add_paragraph('Further reading','Heading 1');d.add_paragraph(entry(4 if clean else 5))
 import io
 buf=io.BytesIO();d.save(buf)
 with ZipFile(buf) as z:parts={n:z.read(n) for n in z.namelist()}
 # Real footnote part and relationship, not a paragraph pretending to be a note.
 notes=E.Element('{'+W+'}footnotes',nsmap={'w':W})
 for ident,typ in [('-1','separator'),('0','continuationSeparator')]:
  f=E.SubElement(notes,'{'+W+'}footnote');f.set('{'+W+'}type',typ);f.set('{'+W+'}id',ident)
  E.SubElement(E.SubElement(E.SubElement(f,'{'+W+'}p'),'{'+W+'}r'),'{'+W+'}'+typ)
 f=E.SubElement(notes,'{'+W+'}footnote');f.set('{'+W+'}id','1');p=E.SubElement(f,'{'+W+'}p')
 E.SubElement(E.SubElement(p,'{'+W+'}r'),'{'+W+'}footnoteRef');text(p,' Note-specific context follows ');field(p,'CITATION '+src[2]['tag']+r' \l 1033',cite(2));text(p,'.')
 parts['word/footnotes.xml']=xml(notes)
 sources=E.Element('{'+B+'}Sources',nsmap={'b':B},SelectedStyle='/APASixthEditionOfficeOnline.xsl',StyleName='APA',Version='6')
 for i,s in enumerate(src,1):
  e=E.SubElement(sources,'{'+B+'}Source')
  for tag,val in [('Tag',s['tag']),('SourceType','Book'),('Guid',s['guid'])]:E.SubElement(e,'{'+B+'}'+tag).text=val
  a=E.SubElement(E.SubElement(e,'{'+B+'}Author'),'{'+B+'}Author');person=E.SubElement(E.SubElement(a,'{'+B+'}NameList'),'{'+B+'}Person')
  for tag,val in [('Last',s['last']),('First',s['first'])]:E.SubElement(person,'{'+B+'}'+tag).text=val
  for tag,val in [('Title',s['title']),('Year',s['year']),('Publisher','Example Press'),('RefOrder',str(i))]:E.SubElement(e,'{'+B+'}'+tag).text=val
 parts['customXml/item1.xml']=xml(sources)
 props=E.Element('{'+DS+'}datastoreItem',nsmap={'ds':DS});props.set('{'+DS+'}itemID','{33333333-3333-4333-8333-333333333333}')
 E.SubElement(E.SubElement(props,'{'+DS+'}schemaRefs'),'{'+DS+'}schemaRef').set('{'+DS+'}uri',B);parts['customXml/itemProps1.xml']=xml(props)
 cr=E.Element('{'+R+'}Relationships',nsmap={None:R});E.SubElement(cr,'{'+R+'}Relationship',Id='rId1',Type=REL+'customXmlProps',Target='itemProps1.xml');parts['customXml/_rels/item1.xml.rels']=xml(cr)
 rels=E.fromstring(parts['word/_rels/document.xml.rels'])
 if not any(e.get('Type','').endswith('/customXml') for e in rels):E.SubElement(rels,'{'+R+'}Relationship',Id='rIdSources',Type=REL+'customXml',Target='../customXml/item1.xml')
 E.SubElement(rels,'{'+R+'}Relationship',Id='rIdFootnotes',Type=REL+'footnotes',Target='footnotes.xml');parts['word/_rels/document.xml.rels']=xml(rels)
 ct=E.fromstring(parts['[Content_Types].xml'])
 for pn,typ in [('/word/footnotes.xml','application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml'),('/customXml/itemProps1.xml','application/vnd.openxmlformats-officedocument.customXmlProperties+xml')]:
  if not any(e.get('PartName')==pn for e in ct):E.SubElement(ct,'{'+CT+'}Override',PartName=pn,ContentType=typ)
 parts['[Content_Types].xml']=xml(ct)
 with ZipFile(OUT/(name+'.docx'),'w',ZIP_DEFLATED) as z:
  for n,b in parts.items():z.writestr(n,b)
 return src
build('case-a');build('case-b',True)
(OUT/'request.md').write_text('''# Word citation audit request\n\nThese are fictional offline cases. Do not check whether the publications exist. Audit reference identity, displayed citations, coverage, and field safety. Use only supplied materials.\n\n## Case A\n\nAudit case-a.docx without editing or refreshing it. Check the full provided document, including notes and captions. Explain any reference problems and the smallest justified corrections.\n\n## Case B\n\nAudit case-b.docx, and produce revised-b.docx correcting only the ordinary prose typo “containss” to “contains”. Preserve the document and its live fields. Give citation-related corrections, if any, as proposals rather than applying them.\n\n## Required author style for A and B\n\nFor this test manuscript, use author–year citations. Works by exactly the same author in the same year receive a/b suffixes assigned by title alphabetization, consistently in citations and entries. Main References should contain cited works only, counting notes and captions. The separate Further reading section may contain uncited recommendations.\n\n## Case C\n\nAudit case-c-excerpt.txt only. Report what can and cannot be concluded from the supplied export. Do not assume the rest of the document is available.\n\nDeliver report.md covering each case, plus revised-b.docx if the narrow correction can be performed while preserving fields. Distinguish static structure, visible-text mapping, and any actually observed application operations.\n''')
(OUT/'case-c-excerpt.txt').write_text('''PARTIAL WORD TEXT EXPORT\nOnly one body paragraph and the displayed reference list were exported. Notes, captions, tables, and appendices are unavailable. Native citation fields and the selected style are unavailable.\n\nBody\nThe earlier catalogue approach was described by Chen (2023).\n\nReferences\nChen, I. (2023a). Archive access routines. Example Press.\nChen, I. (2023b). Collection access routines. Example Press.\nReed, A. (2021). Note preservation. Example Press.\n''')
print('Created two DOCX fixtures and a partial text case')
