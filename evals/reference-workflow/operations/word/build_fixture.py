from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree as E
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json, shutil

OUT=Path(__file__).resolve().parent
B='http://schemas.openxmlformats.org/officeDocument/2006/bibliography'
R='http://schemas.openxmlformats.org/package/2006/relationships'
CT='http://schemas.openxmlformats.org/package/2006/content-types'
DS='http://schemas.openxmlformats.org/officeDocument/2006/customXml'
SOURCES=[
 {'Tag':'Orchid2024','Guid':'{11111111-1111-4111-8111-111111111111}','Last':'Orchid','First':'Ada','Year':'2024','Title':'The Imaginary Copper Orchard','City':'Testville','Publisher':'Fictional Methods Press'},
 {'Tag':'Vale2025','Guid':'{22222222-2222-4222-8222-222222222222}','Last':'Vale','First':'Bram','Year':'2025','Title':'The Invented Paper Observatory','City':'Exampleton','Publisher':'Fictional Methods Press'},
]
def field(p,code,cache):
 for kind,text in [('begin',None),('instruction',code),('separate',None),('text',cache),('end',None)]:
  r=OxmlElement('w:r')
  if kind=='instruction':
   x=OxmlElement('w:instrText'); x.set(qn('xml:space'),'preserve');x.text=' '+text+' '
  elif kind=='text': x=OxmlElement('w:t');x.text=text
  else:
   x=OxmlElement('w:fldChar');x.set(qn('w:fldCharType'),kind)
  r.append(x);p._p.append(r)
doc=Document()
sec=doc.sections[0]
sec.top_margin=sec.bottom_margin=Inches(.85)
sec.left_margin=sec.right_margin=Inches(.9)
for name in ['Normal','Title','Heading 1','Caption']:
 s=doc.styles[name];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
for border in doc.styles.element.xpath('.//w:pBdr'):
 border.getparent().remove(border)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Title'].font.size=Pt(22)
doc.core_properties.author='Fictional Test Author'
doc.core_properties.last_modified_by='Fictional Test Author'
doc.core_properties.title='Native Word citation refresh fixture'
doc.add_paragraph('Native Word citation refresh fixture','Title')
doc.add_paragraph('This fictional document tests native Word source identities and field refresh. Its source records are invented and are not evidence for a research claim.')
p=doc.add_paragraph('Body citation to the orchard source ')
field(p,r'CITATION Orchid2024 \l 1033','[BODY CITATION AWAITING WORD REFRESH]')
p.add_run('.')
doc.add_paragraph('Table 1  Citation outside the body paragraph','Caption')
t=doc.add_table(rows=2,cols=2);t.style='Table Grid'
t.cell(0,0).text='Region';t.cell(0,1).text='Citation'
t.cell(1,0).text='Table cell'
field(t.cell(1,1).paragraphs[0],r'CITATION Vale2025 \l 1033','[TABLE CITATION AWAITING WORD REFRESH]')
doc.add_paragraph('Bibliography','Heading 1')
field(doc.add_paragraph(),r'BIBLIOGRAPHY \l 1033','[BIBLIOGRAPHY AWAITING WORD REFRESH]')
doc.add_paragraph('Fixture boundaries: two source records, two citation occurrences, one bibliography field, no tracked revisions, notes, text boxes, or appendices.')
base=OUT/'fixture-base.docx';doc.save(base)
with ZipFile(base) as z: parts={n:z.read(n) for n in z.namelist()}
src=E.Element('{'+B+'}Sources',nsmap={'b':B},SelectedStyle='/APA.XSL',StyleName='APA',Version='6')
for i,s in enumerate(SOURCES,1):
 x=E.SubElement(src,'{'+B+'}Source')
 for k in ['Tag','SourceType','Guid']:
  E.SubElement(x,'{'+B+'}'+k).text='Book' if k=='SourceType' else s[k]
 a=E.SubElement(E.SubElement(x,'{'+B+'}Author'),'{'+B+'}Author')
 person=E.SubElement(E.SubElement(a,'{'+B+'}NameList'),'{'+B+'}Person')
 for k in ['Last','First']:E.SubElement(person,'{'+B+'}'+k).text=s[k]
 for k in ['Title','Year','City','Publisher']:E.SubElement(x,'{'+B+'}'+k).text=s[k]
 E.SubElement(x,'{'+B+'}RefOrder').text=str(i)
def xml(x):return E.tostring(x,xml_declaration=True,encoding='UTF-8',standalone=True)
parts['customXml/item1.xml']=xml(src)
props=E.Element('{'+DS+'}datastoreItem',nsmap={'ds':DS});props.set('{'+DS+'}itemID','{33333333-3333-4333-8333-333333333333}')
E.SubElement(E.SubElement(props,'{'+DS+'}schemaRefs'),'{'+DS+'}schemaRef').set('{'+DS+'}uri',B)
parts['customXml/itemProps1.xml']=xml(props)
rel=E.Element('{'+R+'}Relationships',nsmap={None:R})
E.SubElement(rel,'{'+R+'}Relationship',Id='rId1',Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXmlProps',Target='itemProps1.xml')
parts['customXml/_rels/item1.xml.rels']=xml(rel)
rels=E.fromstring(parts['word/_rels/document.xml.rels'])
if not any(x.get('Type','').endswith('/customXml') for x in rels):
 E.SubElement(rels,'{'+R+'}Relationship',Id='rIdFixtureSources',Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml',Target='../customXml/item1.xml')
parts['word/_rels/document.xml.rels']=xml(rels)
ct=E.fromstring(parts['[Content_Types].xml'])
if not any(x.get('PartName')=='/customXml/itemProps1.xml' for x in ct):
 E.SubElement(ct,'{'+CT+'}Override',PartName='/customXml/itemProps1.xml',ContentType='application/vnd.openxmlformats-officedocument.customXmlProperties+xml')
parts['[Content_Types].xml']=xml(ct)
with ZipFile(OUT/'before.docx','w',ZIP_DEFLATED) as z:
 for n,data in parts.items():z.writestr(n,data)
shutil.copy2(OUT/'before.docx',OUT/'word-working.docx')
(OUT/'fixture-sources.xml').write_bytes(xml(src))
(OUT/'fixture-spec.json').write_text(json.dumps({'purpose':'Constructed operational probe, not blinded efficacy evaluation','fictional_sources':True,'required_style':'APA as configured for probe; edition is not independently asserted','initial_caches':'Explicit pending Word refresh markers; no simulated post-refresh content','expected_source_tags':['Orchid2024','Vale2025'],'expected_citation_regions':['body','table'],'expected_fields':{'CITATION':2,'BIBLIOGRAPHY':1},'source_records':SOURCES},indent=2)+'\n')
print('Created before.docx and word-working.docx with two fictional native sources and three complex fields.')
